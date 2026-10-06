"""convert.py 테스트. 실행: pytest"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import convert  # noqa: E402


@pytest.fixture
def sample_dir(tmp_path: Path) -> Path:
    src = tmp_path / "src"
    (src / "sub").mkdir(parents=True)

    (src / "표.csv").write_text("이름,점수\n철수,90\n", encoding="utf-8")
    (src / "page.html").write_text("<h1>안녕</h1><p>본문</p>", encoding="utf-8")
    (src / "무시.xyz").write_text("지원하지 않는 확장자", encoding="utf-8")

    from docx import Document
    doc = Document()
    doc.add_heading("워드 제목", level=1)
    doc.add_paragraph("워드 본문입니다.")
    doc.save(src / "sub" / "문서.docx")

    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["품목", "수량"])
    wb.active.append(["사과", 3])
    wb.save(src / "sub" / "재고.xlsx")

    return src


def test_converts_top_level_only_by_default(sample_dir, tmp_path):
    out = tmp_path / "out"
    assert convert.main([str(sample_dir), "-o", str(out)]) == 0

    assert "| 이름 | 점수 |" in (out / "표.md").read_text(encoding="utf-8")
    assert "# 안녕" in (out / "page.md").read_text(encoding="utf-8")
    assert not (out / "무시.md").exists()
    assert not (out / "sub").exists()


def test_recursive_keeps_folder_structure(sample_dir, tmp_path):
    out = tmp_path / "out"
    assert convert.main([str(sample_dir), "-o", str(out), "-r"]) == 0

    assert "# 워드 제목" in (out / "sub" / "문서.md").read_text(encoding="utf-8")
    assert "사과" in (out / "sub" / "재고.md").read_text(encoding="utf-8")


def test_skips_existing_unless_overwrite(sample_dir, tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    (out / "표.md").write_text("기존 내용", encoding="utf-8")
    single = str(sample_dir / "표.csv")

    convert.main([single, "-o", str(out)])
    assert (out / "표.md").read_text(encoding="utf-8") == "기존 내용"

    convert.main([single, "-o", str(out), "--overwrite"])
    assert "철수" in (out / "표.md").read_text(encoding="utf-8")


def test_missing_path_returns_error(tmp_path):
    assert convert.main([str(tmp_path / "없음.pdf"), "-o", str(tmp_path)]) == 1


def test_url_to_filename():
    assert convert.url_to_filename("https://example.com/a/b/") == "example.com_a_b.md"
